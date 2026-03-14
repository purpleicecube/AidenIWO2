#!/usr/bin/env python3
"""Convert markdown slide content to .pptx with Klear.ai brand styling.

Usage:
  python3 md-to-pptx.py --input input.md --output output.pptx [--title "Deck Title"]
  echo "## Slide 1\nContent" | python3 md-to-pptx.py --output output.pptx

Expects markdown with ## headings as slide separators.
Brand treatment derived from Klear.ai masterdeck (Dec 2024 v05).
"""

import argparse
import re
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SCRIPT_DIR = Path(__file__).parent
TEMPLATE_PATH = SCRIPT_DIR / "klearai-template.pptx"
ASSETS_DIR = SCRIPT_DIR / "brand-assets"
LOGO_PATH = ASSETS_DIR / "klearai-logo.png"
LOGO_SM_PATH = ASSETS_DIR / "klearai-logo-sm.png"
HERO_PATH = ASSETS_DIR / "klearai-hero.png"

# Klear.ai brand palette — sourced from masterdeck analysis
NAVY = RGBColor(0x09, 0x1C, 0x53)
DEEP_NAVY = RGBColor(0x04, 0x1F, 0x55)
PURPLE = RGBColor(0x8B, 0x49, 0xE2)
LAVENDER_BG = RGBColor(0xF3, 0xF4, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT = RGBColor(0x2D, 0x2D, 0x2D)
SUBTITLE_GRAY = RGBColor(0x66, 0x66, 0x88)

# Fonts
FONT_HEADING = "Lexend"
FONT_BODY = "Barlow Medium"


def parse_slides(markdown: str) -> tuple[str | None, list[dict]]:
    """Parse markdown into an optional deck title and slide dicts."""
    slides = []
    deck_title = None

    h1_match = re.match(r'^#\s+(.+?)$', markdown.strip(), re.MULTILINE)
    if h1_match:
        deck_title = h1_match.group(1).strip()

    parts = re.split(r'^##\s+', markdown, flags=re.MULTILINE)

    for part in parts:
        part = part.strip()
        if not part:
            continue
        if not slides and not re.match(r'^[^\n]+\n', part):
            continue
        if not slides and h1_match and part.startswith(h1_match.group(1).strip()):
            continue

        lines = part.split('\n', 1)
        title = lines[0].strip().rstrip('#').strip()
        body = lines[1].strip() if len(lines) > 1 else ""

        title = re.sub(r'^Slide\s*\d+\s*[:–-]\s*', '', title, flags=re.IGNORECASE).strip()
        if not title:
            continue

        bullets = []
        for line in body.split('\n'):
            line = line.strip()
            if not line:
                continue
            cleaned = re.sub(r'\*\*([^*]+)\*\*', r'\1', line)
            cleaned = re.sub(r'\*([^*]+)\*', r'\1', cleaned)
            if line.startswith('- ') or line.startswith('* '):
                bullets.append({'text': cleaned[2:].strip(), 'level': 0})
            elif line.startswith('  - ') or line.startswith('  * '):
                bullets.append({'text': cleaned.strip().lstrip('-* ').strip(), 'level': 1})
            else:
                bullets.append({'text': cleaned, 'level': 0})

        slides.append({'title': title, 'bullets': bullets})

    return deck_title, slides


def set_solid_bg(slide, color: RGBColor):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def set_gradient_bg(slide, color_start: RGBColor, color_end: RGBColor):
    bg_elem = slide.background._element
    for old in bg_elem.findall(qn('p:bgPr')):
        bg_elem.remove(old)
    bgPr = etree.SubElement(bg_elem, qn('p:bgPr'))
    gradFill = etree.SubElement(bgPr, qn('a:gradFill'))
    gsLst = etree.SubElement(gradFill, qn('a:gsLst'))
    gs1 = etree.SubElement(gsLst, qn('a:gs'), attrib={'pos': '0'})
    etree.SubElement(gs1, qn('a:srgbClr'), attrib={'val': str(color_start)})
    gs2 = etree.SubElement(gsLst, qn('a:gs'), attrib={'pos': '100000'})
    etree.SubElement(gs2, qn('a:srgbClr'), attrib={'val': str(color_end)})
    etree.SubElement(gradFill, qn('a:lin'), attrib={'ang': '5400000', 'scaled': '1'})
    etree.SubElement(bgPr, qn('a:effectLst'))


def add_rect(slide, left, top, width, height, fill_color, no_line=True):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if no_line:
        shape.line.fill.background()
    return shape


def add_text(slide, left, top, width, height, text, font_name, size_pt, color,
             bold=False, alignment=None, word_wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = word_wrap
    p = tf.paragraphs[0]
    p.clear()
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.bold = bold
    if alignment:
        p.alignment = alignment
    return box


def build_pptx(slides: list[dict], title: str, output_path: str):
    """Build a branded Klear.ai .pptx file."""
    if TEMPLATE_PATH.exists():
        prs = Presentation(str(TEMPLATE_PATH))
    else:
        prs = Presentation()

    sw = prs.slide_width   # 9144000 EMU = 10"
    sh = prs.slide_height  # 5143500 EMU = 5.625"

    has_logo = LOGO_PATH.exists()
    has_logo_sm = LOGO_SM_PATH.exists()
    has_hero = HERO_PATH.exists()

    # ── TITLE SLIDE ──────────────────────────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[9])  # Blank
    set_solid_bg(slide, NAVY)

    # Hero image (right side, behind text — decorative)
    if has_hero:
        hero_w = Inches(5.5)
        hero_h = Inches(4.2)
        slide.shapes.add_picture(
            str(HERO_PATH),
            int(sw - hero_w - Inches(0.2)), Inches(0.7),
            hero_w, hero_h
        )
        # Semi-transparent overlay on hero for text readability
        overlay = add_rect(slide,
            int(sw - hero_w - Inches(0.2)), Inches(0.7),
            hero_w, hero_h,
            NAVY)
        # Set 50% transparency via XML
        spPr = overlay._element.find(qn('p:spPr'))
        if spPr is None:
            spPr = overlay._element.find(qn('a:spPr'))
        fill_elem = spPr.find(qn('a:solidFill'))
        if fill_elem is not None:
            srgb = fill_elem.find(qn('a:srgbClr'))
            if srgb is not None:
                alpha = etree.SubElement(srgb, qn('a:alpha'))
                alpha.set('val', '50000')  # 50% opacity

    # Purple accent bar (top)
    add_rect(slide, 0, 0, sw, Pt(5), PURPLE)

    # Logo (top-left)
    if has_logo:
        slide.shapes.add_picture(str(LOGO_PATH), Inches(0.5), Inches(0.4), Inches(2.2), Inches(0.6))

    # Title text
    add_text(slide, Inches(0.5), Inches(1.8), Inches(6.5), Inches(1.5),
             title, FONT_HEADING, 36, WHITE, bold=True)

    # Subtitle tagline
    add_text(slide, Inches(0.5), Inches(3.3), Inches(5.0), Inches(0.5),
             "AI-Powered Workforce Intelligence", FONT_BODY, 14, PURPLE)

    # Accent line under title
    add_rect(slide, Inches(0.5), Inches(3.15), Inches(2.5), Pt(3), PURPLE)

    # ── CONTENT SLIDES ───────────────────────────────────────────────────
    for i, s in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[9])  # Blank
        set_solid_bg(slide, LAVENDER_BG)

        # Purple accent bar (top)
        add_rect(slide, 0, 0, sw, Pt(5), PURPLE)

        # Purple side accent (left edge, masterdeck style)
        add_rect(slide, 0, Pt(5), Pt(4), int(sh - Pt(5)), PURPLE)

        # Small purple dot (decorative, like masterdeck)
        dot = slide.shapes.add_shape(9, Inches(0.35), Inches(0.55), Inches(0.12), Inches(0.12))  # OVAL
        dot.fill.solid()
        dot.fill.fore_color.rgb = PURPLE
        dot.line.fill.background()

        # Section title
        add_text(slide, Inches(0.6), Inches(0.35), Inches(7.0), Inches(0.55),
                 s['title'].upper(), FONT_HEADING, 22, NAVY, bold=True)

        # Accent line under title
        add_rect(slide, Inches(0.6), Inches(0.95), Inches(3.5), Pt(2), PURPLE)

        # Bullets
        bullet_top = Inches(1.2)
        bullet_box = slide.shapes.add_textbox(
            Inches(0.6), bullet_top, Inches(8.5), int(sh - bullet_top - Inches(0.6))
        )
        tf = bullet_box.text_frame
        tf.word_wrap = True

        for j, bullet in enumerate(s['bullets']):
            if j == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()

            p.clear()
            run = p.add_run()
            p.space_after = Pt(6)

            if bullet['level'] == 0:
                run.text = f"\u2022  {bullet['text']}"
                run.font.name = FONT_BODY
                run.font.size = Pt(14)
                run.font.color.rgb = DEEP_NAVY
            else:
                run.text = f"    \u25E6  {bullet['text']}"
                run.font.name = FONT_BODY
                run.font.size = Pt(12)
                run.font.color.rgb = DARK_TEXT

        # Small logo (bottom-right)
        if has_logo_sm:
            slide.shapes.add_picture(
                str(LOGO_SM_PATH),
                int(sw - Inches(2.0)), int(sh - Inches(0.55)),
                Inches(1.3), Inches(0.37)
            )

        # Slide number (bottom-left)
        add_text(slide, Inches(0.5), int(sh - Inches(0.45)), Inches(0.5), Inches(0.3),
                 str(i + 1), FONT_BODY, 9, SUBTITLE_GRAY, alignment=PP_ALIGN.LEFT)

    # ── CLOSING SLIDE ────────────────────────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[9])  # Blank
    set_gradient_bg(slide, PURPLE, NAVY)

    # White accent bar (top)
    add_rect(slide, 0, 0, sw, Pt(4), WHITE)

    # Logo (center-ish)
    if has_logo:
        slide.shapes.add_picture(
            str(LOGO_PATH),
            int(sw / 2 - Inches(1.5)), Inches(1.2),
            Inches(3.0), Inches(0.85)
        )

    # Thank You text
    add_text(slide, Inches(1), Inches(2.5), Inches(8), Inches(1.2),
             "Thank You", FONT_HEADING, 44, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    # Subtitle
    add_text(slide, Inches(1), Inches(3.6), Inches(8), Inches(0.5),
             "klear.ai", FONT_BODY, 16, LAVENDER_BG, alignment=PP_ALIGN.CENTER)

    # Accent line
    add_rect(slide, int(sw / 2 - Inches(1.5)), Inches(3.45), Inches(3.0), Pt(3), WHITE)

    prs.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Convert markdown slides to .pptx")
    parser.add_argument('--input', '-i', help='Input markdown file (or stdin)')
    parser.add_argument('--output', '-o', required=True, help='Output .pptx file path')
    parser.add_argument('--title', '-t', default='Presentation', help='Deck title for title slide')
    args = parser.parse_args()

    if args.input:
        markdown = Path(args.input).read_text()
    else:
        markdown = sys.stdin.read()

    if not markdown.strip():
        print("Error: empty input", file=sys.stderr)
        sys.exit(1)

    deck_title, slides = parse_slides(markdown)
    if not slides:
        print("Error: no slides found (expected ## headings)", file=sys.stderr)
        sys.exit(1)

    final_title = args.title if args.title != 'Presentation' else (deck_title or args.title)

    output = build_pptx(slides, final_title, args.output)
    import json
    print(json.dumps({
        "success": True,
        "path": str(output),
        "slideCount": len(slides) + 2,
    }))


if __name__ == '__main__':
    main()
